"""Export the manuscript markdown sources to submission DOCX files.

Regenerates the four manuscript DOCX files from the current .md sources so the
exports can never drift from the sources that the placeholder scanner audits.
No third-party measurement data is read or written by this script.

Usage:
    python scripts/export_manuscript_docx.py
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = ROOT / "manuscript"

FILES = [
    "Bioinformatics_Main_Final_EN.md",
    "Bioinformatics_Main_Final_CN.md",
    "Bioinformatics_Supplementary_Final_EN.md",
    "Bioinformatics_Supplementary_Final_CN.md",
]

BOLD_ITALIC_RE = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*.+?\*|`.+?`)")


def add_runs(paragraph, text: str) -> None:
    """Render **bold**, *italic*, ***both*** and `code` spans into runs."""
    for part in BOLD_ITALIC_RE.split(text):
        if not part:
            continue
        if part.startswith("***") and part.endswith("***"):
            run = paragraph.add_run(part[3:-3])
            run.bold = True
            run.italic = True
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            paragraph.add_run(part)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    """Append one markdown table as a DOCX table with a bold header row."""
    # Drop the markdown header-separator row (e.g. |---|---|) before layout.
    rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c.strip()) for c in r)]
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            para = table.cell(i, j).paragraphs[0]
            add_runs(para, cell.strip())
            for run in para.runs:
                run.font.size = Pt(9)
                if i == 0:
                    run.bold = True


def convert(md_path: Path, docx_path: Path) -> None:
    """Convert one manuscript markdown file to a DOCX file with basic block rendering."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]

        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c for c in lines[i].strip().strip("|").split("|")]
                rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue

        if re.fullmatch(r"-{3,}", line.strip()):
            doc.add_paragraph()
            i += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            if level == 1:
                para = doc.add_heading("", level=0)
            else:
                para = doc.add_heading("", level=level - 1)
            add_runs(para, heading.group(2))
            i += 1
            continue

        if re.match(r"^\s*[-*]\s+", line):
            para = doc.add_paragraph(style="List Bullet")
            add_runs(para, re.sub(r"^\s*[-*]\s+", "", line))
            i += 1
            continue

        if re.match(r"^\s*\d+\.\s+", line):
            para = doc.add_paragraph(style="List Number")
            add_runs(para, re.sub(r"^\s*\d+\.\s+", "", line))
            i += 1
            continue

        if line.startswith(">"):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Pt(24)
            add_runs(para, line.lstrip("> "))
            i += 1
            continue

        if line.strip():
            para = doc.add_paragraph()
            add_runs(para, line.strip())
        i += 1

    doc.save(docx_path)
    print(f"wrote {docx_path.relative_to(ROOT)}")


def main() -> None:
    for name in FILES:
        md_path = MANUSCRIPT / name
        if not md_path.exists():
            raise SystemExit(f"missing source: {md_path}")
        convert(md_path, md_path.with_suffix(".docx"))


if __name__ == "__main__":
    main()
