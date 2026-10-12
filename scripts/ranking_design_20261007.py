"""Read-only workbook structure audit and label-free rank-null simulation."""
import hashlib
import json
import math
from pathlib import Path
import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'artifacts/external/ranking-design-20261007'


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def main():
    OUT.mkdir(exist_ok=True)
    p = ROOT/'artifacts/external/temporal-absolute-kinetics/edmond-S0HJ48/raw/Kinetic_Acomosus_and_variant.xlsx'
    original_hash = sha(p)
    w = openpyxl.load_workbook(p, read_only=True, data_only=False)
    s = w['Tabelle1']
    # Report observed layout, never assume columns are independent biological replicates.
    blocks = []
    for label, concentration, start, end in [('acetylCoA_WT','A',2,3), ('acetylCoA_deletion','A',4,6),
                                           ('oxaloacetate_WT','J',11,12), ('oxaloacetate_deletion','J',13,14)]:
        count = [sum(s.cell(row,c).value is not None for c in range(start,end+1)) for row in range(4,11)]
        blocks.append(dict(block=label, concentration_range=f'{concentration}4:{concentration}10',
            measurement_range=f'{s.cell(4,start).coordinate}:{s.cell(10,end).coordinate}',
            visible_measurement_columns_per_concentration=count,
            concentration_uM=[s[f'{concentration}{i}'].value for i in range(4,11)]))
    formulas = [c.coordinate for row in s for c in row if c.data_type=='f']
    w.close()
    assert sha(p)==original_hash
    source = dict(workbook=str(p.relative_to(ROOT)), workbook_sha256=original_hash, sheet='Tabelle1', blocks=blocks,
        formulas_found=formulas, workbook_modified=False, raw_time_traces_recovered=False,
        fixed_cosubstrate_evidence='Main Kinetic enzyme assays: oxaloacetate 1 mM for varied acetyl-CoA; acetyl-CoA 0.5 mM for varied oxaloacetate.',
        fixed_conditions_resolved=True, pending=['SI Figure12 says n=3 technical replicates; visible workbook columns are 2/3/2/2, relation unresolved.',
        'Workbook labels substrate concentration units but not the response-value units; do not infer or refit.',
        'Published rate/parameter errors and biological independence cannot be reconstructed from this layout alone.'],
        candidate_ids=['cs-edmond-017','cs-edmond-018','cs-edmond-019','cs-edmond-020'],
        ranking_status='hold_repeat_mapping_and_response_units', absolute_status_changed=False,
        primary_xml_sha256=sha(OUT/'PMC11615384.xml'))
    (OUT/'citrate-workbook-audit.json').write_text(json.dumps(source,indent=2)+'\n')
    review_path=ROOT/'artifacts/external/ranking-campaign-review-20261006/review.json'
    review=json.loads(review_path.read_text())
    groups=[g for g in review['groups'] if g['status']=='conditions_supported_candidate']
    sensitivity=[]
    for n in [2,3,5,8]:
        kept=[g for g in groups if g['distinct_sequence_signatures']>=n]
        sensitivity.append(dict(minimum_candidates=n,groups=len(kept),records=sum(g['records'] for g in kept),sources=len({g['source_id'] for g in kept})))
    rng=np.random.default_rng(20261007)
    trials=20000
    simulation=[]
    for n in [2,3,4,5,8,10,21]:
        fixed=np.arange(n,dtype=float);fixed-=fixed.mean()
        ranks=np.argsort(rng.random((trials,n)),axis=1).astype(float);ranks-=ranks.mean(axis=1,keepdims=True)
        rho=(ranks@fixed)/(fixed@fixed)
        simulation.append(dict(n=n,trials=trials,mean_rho=float(rho.mean()),sd_rho=float(rho.std(ddof=1)),
            probability_abs_rho_at_least_0_8=float(np.mean(np.abs(rho)>=.8-1e-12)),
            exact_two_sided_probability_perfect_order_under_no_ties=2/math.factorial(n)))
    result=dict(date='2026-10-07', seed=20261007,simulation='independent uniform random ranks; no ties; NOT power or model performance',
        external_labels_read=False, external_predictions_read=False, review_sha256=sha(review_path),
        metadata_size_sensitivity=sensitivity, null_simulations=simulation,
        minimum_size_frozen=False, significance_gate_created=False)
    (OUT/'design-simulation.json').write_text(json.dumps(result,indent=2)+'\n')
    print(json.dumps(dict(workbook_visible_columns=[x['visible_measurement_columns_per_concentration'][0] for x in blocks],
                         formulas=formulas,metadata_size_sensitivity=sensitivity,simulations=simulation),indent=2))


if __name__=='__main__':
    main()
